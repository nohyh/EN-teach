"""教材文件抽取、启发式组件生成与课程契约校验。"""
from __future__ import annotations

import io
import json
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

from pydantic import ValidationError

from app.schemas.lesson import CourseContent


class ContentParseError(ValueError):
    def __init__(self, message: str, *, location: str = "file") -> None:
        super().__init__(message)
        self.location = location


def _validation_errors(error: ValidationError) -> list[dict[str, str]]:
    return [
        {
            "location": ".".join(str(part) for part in item["loc"]),
            "message": item["msg"],
            "type": item["type"],
        }
        for item in error.errors()
    ]


def validate_course_content(value: Any) -> tuple[dict[str, Any] | None, list[dict[str, str]]]:
    try:
        content = CourseContent.model_validate(value)
    except ValidationError as error:
        return None, _validation_errors(error)

    errors: list[dict[str, str]] = []
    section_ids: set[str] = set()
    for section_index, section in enumerate(content.sections):
        if section.id in section_ids:
            errors.append({
                "location": f"sections.{section_index}.id",
                "message": "小节 ID 重复",
                "type": "duplicate_section_id",
            })
        section_ids.add(section.id)
        activity_keys: set[str] = set()
        for activity_index, activity in enumerate(section.activities):
            payload = activity.model_dump(mode="json")
            key = json.dumps(payload, ensure_ascii=False, sort_keys=True)
            if key in activity_keys:
                errors.append({
                    "location": f"sections.{section_index}.activities.{activity_index}",
                    "message": "小节中存在完全重复的活动",
                    "type": "duplicate_activity",
                })
            activity_keys.add(key)
            if payload["type"] == "recall" and payload["mode"] == "fill_blank" and "____" not in payload["prompt"]:
                errors.append({
                    "location": f"sections.{section_index}.activities.{activity_index}.prompt",
                    "message": "填空题 prompt 必须包含 ____",
                    "type": "missing_blank",
                })
    return (content.model_dump(mode="json") if not errors else None), errors


def _slug(value: str, fallback: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return (slug or fallback)[:80]


def _normalize_json(value: Any, title_hint: str) -> dict[str, Any]:
    if isinstance(value, list):
        value = {"title": title_hint, "intro": "", "activities": value}
    if not isinstance(value, dict):
        raise ContentParseError("JSON 顶层必须是对象或活动数组", location="json")
    if "sections" in value:
        return value
    activities = value.get("activities")
    if not isinstance(activities, list):
        raise ContentParseError("JSON 缺少 sections 或 activities 数组", location="json.activities")

    groups: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for index, raw in enumerate(activities):
        if not isinstance(raw, dict):
            raise ContentParseError("活动必须是对象", location=f"json.activities.{index}")
        section_id = str(raw.get("sectionId") or "section-1")
        section_title = str(raw.get("sectionTitle") or "第一课")
        component = {key: val for key, val in raw.items() if key not in {"sectionId", "sectionTitle"}}
        group = groups.setdefault(section_id, {"id": section_id, "title": section_title, "activities": []})
        group["activities"].append(component)
    return {
        "title": str(value.get("title") or title_hint),
        "intro": str(value.get("intro") or ""),
        "sections": list(groups.values()),
    }


def _extract_docx(payload: bytes) -> str:
    try:
        from docx import Document

        document = Document(io.BytesIO(payload))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(lines)
    except Exception as error:
        raise ContentParseError("DOCX 无法读取或文件已损坏", location="docx") from error


def _extract_pdf(payload: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(payload))
    except Exception as error:
        raise ContentParseError("PDF 无法读取或文件已加密/损坏", location="pdf") from error
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception as error:
            raise ContentParseError("PDF 页面文本抽取失败", location=f"pdf.page.{index}") from error
    text = "\n".join(pages).strip()
    if not text:
        raise ContentParseError("PDF 没有可提取文本，可能是扫描件，请上传页面图片进入 OCR 队列", location="pdf")
    return text


def _text_to_course(text: str, title_hint: str) -> dict[str, Any]:
    current_id = "section-1"
    current_title = "第一课"
    heading_count = 0
    heading_ids: set[str] = set()
    sections: OrderedDict[str, dict[str, Any]] = OrderedDict()
    pairs: list[tuple[str, str, str, str]] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^#{1,6}\s+(.+)$", line)
        if heading:
            heading_count += 1
            current_title = heading.group(1).strip()
            candidate = _slug(current_title, f"section-{heading_count}")
            current_id = candidate if candidate not in heading_ids else f"{candidate}-{heading_count}"
            heading_ids.add(current_id)
            continue
        parts = re.split(r"\s*(?:\||\t|\s[-—–:]\s)\s*", re.sub(r"^[-*+]\s+", "", line), maxsplit=1)
        if len(parts) != 2:
            continue
        english, chinese = parts[0].strip(), parts[1].strip()
        if re.search(r"[A-Za-z]", chinese) and re.search(r"[\u4e00-\u9fff]", english):
            english, chinese = chinese, english
        if not re.search(r"[A-Za-z]", english) or not re.search(r"[\u4e00-\u9fff]", chinese):
            continue
        pairs.append((current_id, current_title, english, chinese))

    if not pairs:
        raise ContentParseError(
            "未识别到英中内容对；请使用“apple - 苹果”或“I like apples. | 我喜欢苹果。”格式",
            location="text.lines",
        )

    for section_id, section_title, english, chinese in pairs:
        section = sections.setdefault(section_id, {"id": section_id, "title": section_title, "activities": []})
        is_word = re.fullmatch(r"[A-Za-z][A-Za-z'’-]*", english) is not None
        if is_word:
            example = f"This is {english}."
            example_meaning = f"这是{chinese}。"
            section["activities"].extend([
                {"type": "word", "word": english, "meaning": chinese, "example": example, "exampleMeaning": example_meaning, "message": f"一起学习 {english}！"},
                {"type": "recall", "mode": "en_to_zh", "prompt": english, "answer": chinese, "message": f"{english} 是什么意思？"},
                {"type": "recall", "mode": "zh_to_en", "prompt": chinese, "answer": english, "message": f"{chinese}用英语怎么说？"},
                {"type": "pronunciation", "content": english, "meaning": chinese, "message": f"跟着 Lumi 读：{english}"},
            ])
        else:
            section["activities"].extend([
                {"type": "sentence", "sentence": english, "meaning": chinese, "message": "一起学习这个句子。"},
                {"type": "recall", "mode": "zh_to_en", "prompt": chinese, "answer": english, "message": "试着说出对应的英文句子。"},
                {"type": "pronunciation", "content": english, "meaning": chinese, "message": "跟着 Lumi 读完整句子。"},
            ])

    for section in sections.values():
        section["activities"].append({
            "type": "dialog",
            "scene": section["title"],
            "goal": "使用本节学到的单词和句子完成简短英语对话",
            "opening": "Hello! What did you learn today?",
        })
    return {"title": title_hint, "intro": "由上传教材自动生成的候选课程，发布前请人工审核。", "sections": list(sections.values())}


def parse_source(filename: str, payload: bytes) -> tuple[dict[str, Any], str]:
    suffix = Path(filename).suffix.lower()
    title_hint = Path(filename).stem[:255] or "未命名课程"
    if suffix == ".json":
        try:
            text = payload.decode("utf-8-sig")
            return _normalize_json(json.loads(text), title_hint), text
        except UnicodeDecodeError as error:
            raise ContentParseError("JSON 必须使用 UTF-8 编码", location="json") from error
        except json.JSONDecodeError as error:
            raise ContentParseError(error.msg, location=f"json.line.{error.lineno}.column.{error.colno}") from error
    if suffix in {".md", ".markdown", ".txt"}:
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ContentParseError("文本文件必须使用 UTF-8 编码", location="text") from error
    elif suffix == ".docx":
        text = _extract_docx(payload)
    elif suffix == ".pdf":
        text = _extract_pdf(payload)
    else:
        raise ContentParseError(f"不支持解析 {suffix or '无扩展名'} 文件", location="file.extension")
    return _text_to_course(text, title_hint), text
